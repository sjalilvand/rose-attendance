"""HR monthly Excel export tests — report engine + API (فاز ۲، گام ۶).

Pins the workbook contract: three sheets, right-to-left views, per-(person ×
branch) summary rows read from MonthlySummary, the LeaveLedger snapshot and
the HR-only access guard of both endpoints.
"""
from __future__ import annotations

import datetime as dt
import io
import uuid

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.attendance.models import AttendanceEvent
from apps.attendance.services import apply_event_to_pairing
from apps.core.enums import RestMode
from apps.leaves.models import LeaveType
from apps.orgs.models import Branch, Region
from apps.people.models import Employee, Role, RoleAssignment, User
from apps.scheduling.models import ScheduledShift, Shift

from .models import MonthlySummary
from .services import recompute_month_summary

from .excel_export import (
    SHEET_BALANCES,
    SHEET_GUIDE,
    SHEET_SUMMARY,
    build_monthly_report,
    month_preview,
)


def make_branch(code="B100"):
    region, _ = Region.objects.get_or_create(code="R1", defaults={"name": "منطقه یک"})
    return Branch.objects.create(
        code=code, name=f"شعبه {code}", region=region,
        geo_latitude=35.7997, geo_longitude=51.4344,
        geo_radius_m=150, city="تهران",
    )


def make_employee(branch, code="12345", first="علی", last="رضایی"):
    emp = Employee.objects.create(
        personnel_code=code, first_name=first, last_name=last, home_branch=branch,
    )
    user = User.objects.create_user(
        username=f"u{code}", password="S3cret!Pass", employee=emp
    )
    return emp, user


def grant(user, code, branch=None):
    RoleAssignment.objects.create(
        user=user, role=Role.objects.get(code=code), branch=branch,
        valid_from=timezone.localdate() - dt.timedelta(days=1),
    )


def tehran(date_: dt.date, hour: int, minute: int = 0):
    return dt.datetime.combine(date_, dt.time(hour, minute)).replace(
        tzinfo=timezone.get_current_timezone()
    )


class ExcelReportBase(TestCase):
    """Month 1405/07 with one worked day (9:15→17:00) — mirrors RollupTestBase."""

    def setUp(self):
        self.branch = make_branch()
        self.emp, self.user = make_employee(self.branch)
        self.hr_emp, self.hr_user = make_employee(self.branch, code="HR-9", first="مریم", last="مدیر")
        grant(self.hr_user, "HR")
        self.shift = Shift.objects.create(
            code="D1", name="اداری", starts_at=dt.time(9, 0), ends_at=dt.time(17, 0),
            rest_minutes=60, rest_mode=RestMode.FIXED,
        )
        self.work_date = dt.date(2026, 10, 12)  # ۲۰ مهر ۱۴۰۵
        ScheduledShift.objects.create(
            employee=self.emp, branch=self.branch, work_date=self.work_date, shift=self.shift,
        )
        for kind, when in (("CHECK_IN", tehran(self.work_date, 9, 15)),
                           ("CHECK_OUT", tehran(self.work_date, 17, 0))):
            event = AttendanceEvent.objects.create(
                event_uuid=uuid.uuid4(), employee=self.emp, branch=self.branch,
                event_type=kind, device_time=when, effective_time=when,
                channel="PWA", connection_mode="OFFLINE_SYNC", event_status="ACCEPTED",
            )
            apply_event_to_pairing(event)
        self.summary = recompute_month_summary(self.emp.id, self.branch.id, 1405, 7)
        self.client = APIClient()


class WorkbookContentTests(ExcelReportBase):
    def test_three_sheets_and_summary_row(self):
        payload = build_monthly_report(1405, 7)
        wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        self.assertEqual(wb.sheetnames, [SHEET_SUMMARY, SHEET_BALANCES, SHEET_GUIDE])
        ws = wb[SHEET_SUMMARY]
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "کد پرسنلی")
        data = [r for r in rows[1:] if r[0] is not None]
        self.assertEqual(len(data), 1)
        row = data[0]
        self.assertEqual(row[0], "12345")
        self.assertEqual(row[2], "B100")
        self.assertEqual(row[4], 420)   # دقیقه لازم (480 شیفت − 60 استراحت)
        self.assertEqual(row[5], 405)   # کار پذیرفته
        self.assertEqual(row[7], 15)    # تأخیر
        self.assertEqual(row[-1], "پیش‌نویس")
        wb.close()

    def test_balance_sheet_reflects_ledger_sum(self):
        from apps.leaves.models import LeaveLedger

        annual, _ = LeaveType.objects.get_or_create(code="ANNUAL", defaults={"name": "استحقاقی"})
        LeaveLedger.objects.create(
            employee=self.emp, leave_type=annual, entry_kind="OPENING",
            minutes=12480, balance_after_minutes=12480,
            basis_date=dt.date(2026, 3, 21),
        )
        payload = build_monthly_report(1405, 7)
        wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        rows = list(wb[SHEET_BALANCES].iter_rows(values_only=True))
        self.assertEqual(rows[0][4], "مانده (دقیقه)")
        data = [r for r in rows[1:] if r[0] is not None]
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0][0], "12345")
        self.assertEqual(data[0][3], "ANNUAL")
        self.assertEqual(data[0][4], 12480)
        self.assertEqual(data[0][5], "208:00")  # 12480 دقیقه = ۲۰۸ ساعت
        wb.close()

    def test_leave_type_columns_appear_when_used(self):
        MonthlySummary.objects.filter(pk=self.summary.pk).update(
            leave_minutes_by_type={"ANNUAL": 240, "SICK": 120}
        )
        payload = build_monthly_report(1405, 7)
        wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        header = list(wb[SHEET_SUMMARY].iter_rows(values_only=True))[0]
        self.assertIn("مرخصی ANNUAL", header)
        self.assertIn("مرخصی SICK", header)
        wb.close()

    def test_preview_totals(self):
        data = month_preview(1405, 7)
        self.assertEqual(data["rows"], 1)
        self.assertEqual(data["totals"]["employees"], 1)
        self.assertEqual(data["totals"]["required_minutes"], 420)
        self.assertEqual(data["by_branch"][0]["branch"], "B100")
        self.assertEqual(data["by_branch"][0]["required_minutes"], 420)


class ReportApiTests(ExcelReportBase):
    def test_hr_can_download_excel(self):
        self.client.force_authenticate(self.hr_user)
        res = self.client.get("/api/v1/reports/monthly/excel/?year=1405&month=7")
        self.assertEqual(res.status_code, 200)
        self.assertIn("spreadsheetml", res["Content-Type"])
        self.assertIn("rose-monthly-report-1405-07.xlsx", res["Content-Disposition"])
        wb = load_workbook(io.BytesIO(res.content), read_only=True)
        self.assertEqual(wb.sheetnames[0], SHEET_SUMMARY)
        wb.close()

    def test_regular_employee_is_forbidden(self):
        self.client.force_authenticate(self.user)
        res = self.client.get("/api/v1/reports/monthly/excel/?year=1405&month=7")
        self.assertEqual(res.status_code, 403)
        res2 = self.client.get("/api/v1/reports/monthly/preview/?year=1405&month=7")
        self.assertEqual(res2.status_code, 403)

    def test_invalid_month_is_400(self):
        self.client.force_authenticate(self.hr_user)
        res = self.client.get("/api/v1/reports/monthly/excel/?year=1405&month=13")
        self.assertEqual(res.status_code, 400)

    def test_unknown_branch_is_404(self):
        self.client.force_authenticate(self.hr_user)
        res = self.client.get("/api/v1/reports/monthly/preview/?year=1405&month=7&branch=NOPE")
        self.assertEqual(res.status_code, 404)

    def test_preview_with_refresh_recomputes(self):
        # جلسات ناقص بعد از rollup اولیه اضافه شد → refresh باید آن را ببیند
        self.client.force_authenticate(self.hr_user)
        ScheduledShift.objects.create(
            employee=self.emp, branch=self.branch,
            work_date=self.work_date + dt.timedelta(days=1), shift=self.shift,
        )
        res = self.client.get("/api/v1/reports/monthly/preview/?year=1405&month=7&refresh=0")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data["totals"]["required_minutes"], 420)
        res2 = self.client.get("/api/v1/reports/monthly/preview/?year=1405&month=7")
        self.assertEqual(res2.status_code, 200)
        self.assertEqual(res2.data["totals"]["required_minutes"], 840)
