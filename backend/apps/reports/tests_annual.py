"""Annual + disciplinary report tests — step 8 of phase 2.

Pins:
- late_days / absence_days persistence in the repeatable rollup
- year_preview trend aggregation (per-month totals of persisted rows)
- annual workbook contract (روند ماهانه / خلاصه سالانه پرسنل / راهنما)
- disciplinary workbook contract (انتظامی تأخیر/غیبت / راهنما)
- HR-only access + refresh semantics of the three new endpoints
"""
from __future__ import annotations

import datetime as dt
import io

from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.scheduling.models import ScheduledShift

from .excel_export import (
    SHEET_DISCIPLINARY,
    SHEET_GUIDE,
    SHEET_YEAR_PEOPLE,
    SHEET_YEAR_TREND,
    build_annual_report,
    build_disciplinary_report,
    year_preview,
)
from .services import recompute_month_summary

from .tests_excel import ExcelReportBase


class RollupCountersTests(ExcelReportBase):
    def test_late_day_counted_once(self):
        self.summary.refresh_from_db()
        self.assertEqual(self.summary.late_minutes, 15)
        self.assertEqual(self.summary.late_days, 1)
        self.assertEqual(self.summary.absence_days, 0)

    def test_absent_scheduled_day_counts(self):
        ScheduledShift.objects.create(
            employee=self.emp, branch=self.branch,
            work_date=self.work_date + dt.timedelta(days=1), shift=self.shift,
        )
        summary = recompute_month_summary(self.emp.id, self.branch.id, 1405, 7)
        self.assertEqual(summary.absence_days, 1)
        self.assertEqual(summary.late_days, 1)


class YearPreviewTests(ExcelReportBase):
    def test_two_months_aggregated(self):
        ScheduledShift.objects.create(
            employee=self.emp, branch=self.branch,
            work_date=dt.date(2026, 11, 2), shift=self.shift,  # ۱۱ آبان ۱۴۰۵
        )
        recompute_month_summary(self.emp.id, self.branch.id, 1405, 8)
        data = year_preview(1405)
        self.assertEqual(data["period_year"], 1405)
        months = {m["jm"]: m for m in data["months"]}
        self.assertEqual(set(months), {7, 8})
        self.assertEqual(months[7]["late_minutes"], 15)
        self.assertEqual(months[8]["employees"], 1)
        self.assertEqual(data["totals"]["late_minutes"], 15)
        self.assertEqual(months[7]["month_name"], "مهر")

    def test_branch_scoping(self):
        recompute_month_summary(self.emp.id, self.branch.id, 1405, 8)
        data = year_preview(1405, branch=self.branch)
        self.assertEqual(data["branch"], "B100")
        self.assertEqual(data["rows"], 2)


class AnnualWorkbookTests(ExcelReportBase):
    def test_sheets_and_trend_row(self):
        payload = build_annual_report(1405)
        wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        self.assertEqual(wb.sheetnames, [SHEET_YEAR_TREND, SHEET_YEAR_PEOPLE, SHEET_GUIDE])
        trend = [r for r in wb[SHEET_YEAR_TREND].iter_rows(values_only=True)][1:]
        trend = [r for r in trend if r[0]]
        self.assertEqual(len(trend), 1)
        row = trend[0]
        self.assertIn("مهر", row[0])
        self.assertEqual(row[1], 1)     # پرسنل
        self.assertEqual(row[2], 420)   # دقیقه لازم
        self.assertEqual(row[3], 405)   # کار پذیرفته
        self.assertEqual(row[4], 15)    # تأخیر
        self.assertEqual(row[5], 1)     # روز تأخیر
        people = [r for r in wb[SHEET_YEAR_PEOPLE].iter_rows(values_only=True)][1:]
        people = [r for r in people if r[0]]
        self.assertEqual(len(people), 1)
        self.assertEqual(people[0][0], "12345")
        self.assertEqual(people[0][3], 1)    # ماه‌های دارای رکورد
        self.assertEqual(people[0][7], 1)    # جمع روز تأخیر
        wb.close()


class DisciplinaryWorkbookTests(ExcelReportBase):
    def test_sheets_and_row(self):
        payload = build_disciplinary_report(1405, 7)
        wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        self.assertEqual(wb.sheetnames, [SHEET_DISCIPLINARY, SHEET_GUIDE])
        rows = list(wb[SHEET_DISCIPLINARY].iter_rows(values_only=True))
        self.assertEqual(rows[0][4], "تعداد روز تأخیر")
        data = [r for r in rows[1:] if r[0] is not None]
        self.assertEqual(len(data), 1)
        row = data[0]
        self.assertEqual(row[0], "12345")
        self.assertEqual(row[4], 1)    # روز تأخیر
        self.assertEqual(row[5], 15)   # جمع تأخیر
        self.assertEqual(row[6], 0)    # روز غیبت
        self.assertEqual(row[8], 15)   # کسری (لازم ۴۲۰ − کار ۴۰۵)
        wb.close()


class NewReportApiTests(ExcelReportBase):
    def setUp(self):
        super().setUp()
        self.client.force_authenticate(self.hr_user)

    def test_annual_preview_and_excel(self):
        res = self.client.get("/api/v1/reports/annual/preview/?year=1405&refresh=0")
        self.assertEqual(res.status_code, res and 200, res.data if hasattr(res, "data") else res)
        self.assertEqual(len(res.data["months"]), 1)
        res2 = self.client.get("/api/v1/reports/annual/excel/?year=1405&refresh=0")
        self.assertEqual(res2.status_code, 200)
        self.assertIn("rose-annual-report-1405.xlsx", res2["Content-Disposition"])
        wb = load_workbook(io.BytesIO(res2.content), read_only=True)
        self.assertEqual(wb.sheetnames[0], SHEET_YEAR_TREND)
        wb.close()

    def test_annual_refresh_recomputes_missing_month(self):
        ScheduledShift.objects.create(
            employee=self.emp, branch=self.branch,
            work_date=dt.date(2026, 11, 3), shift=self.shift,
        )
        res = self.client.get("/api/v1/reports/annual/preview/?year=1405&refresh=0")
        self.assertEqual(len(res.data["months"]), 1)
        res2 = self.client.get("/api/v1/reports/annual/preview/?year=1405&refresh=1")
        self.assertEqual(res2.status_code, 200)
        self.assertEqual(len(res2.data["months"]), 2)

    def test_disciplinary_excel_download(self):
        res = self.client.get("/api/v1/reports/disciplinary/excel/?year=1405&month=7")
        self.assertEqual(res.status_code, 200)
        self.assertIn("rose-disciplinary-1405-07.xlsx", res["Content-Disposition"])
        wb = load_workbook(io.BytesIO(res.content), read_only=True)
        self.assertEqual(wb.sheetnames[0], SHEET_DISCIPLINARY)
        wb.close()

    def test_employee_forbidden_on_new_endpoints(self):
        self.client.force_authenticate(self.user)
        for url in (
            "/api/v1/reports/annual/preview/?year=1405",
            "/api/v1/reports/annual/excel/?year=1405",
            "/api/v1/reports/disciplinary/excel/?year=1405&month=7",
        ):
            self.assertEqual(self.client.get(url).status_code, 403, url)

    def test_invalid_year_is_400(self):
        self.client.force_authenticate(self.hr_user)
        res = self.client.get("/api/v1/reports/annual/preview/?year=99")
        self.assertEqual(res.status_code, 400)

    def test_monthly_preview_reports_day_counts(self):
        res = self.client.get("/api/v1/reports/monthly/preview/?year=1405&month=7&refresh=0")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data["totals"]["late_days"], 1)
        self.assertEqual(res.data["totals"]["absence_days"], 0)
