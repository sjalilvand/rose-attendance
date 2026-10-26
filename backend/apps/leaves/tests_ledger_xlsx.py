"""Leave balance Excel import tests — dedicated template + importer (گام ۶).

Pins: template carries live leave-type reference, importer accepts Persian
headers and Jalali dates, dry run touches nothing, good rows land while bad
rows are reported, negatives pass only as ADJUST, and both endpoints stay
HR-only.
"""
from __future__ import annotations

import datetime as dt
import io

from django.core.cache import cache
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.leaves.models import LeaveLedger, LeaveType
from apps.orgs.models import Branch, Region
from apps.people.models import Employee, Role, RoleAssignment, User

from .ledger import (
    BALANCE_HEADER_ALIASES,
    WorkbookError,
    build_balance_template,
    current_balance_minutes,
    import_balances,
)


def make_branch(code="B100"):
    region, _ = Region.objects.get_or_create(code="R1", defaults={"name": "منطقه یک"})
    return Branch.objects.create(
        code=code, name=f"شعبه {code}", region=region,
        geo_latitude=35.7997, geo_longitude=51.4344, geo_radius_m=150,
    )


def make_employee(branch, code="12345"):
    emp = Employee.objects.create(
        personnel_code=code, first_name="علی", last_name="رضایی", home_branch=branch,
    )
    user = User.objects.create_user(username=f"u{code}", password="S3cret!Pass", employee=emp)
    return emp, user


def grant(user, code, branch=None):
    RoleAssignment.objects.create(
        user=user, role=Role.objects.get(code=code), branch=branch,
        valid_from=timezone.localdate() - dt.timedelta(days=1),
    )


def leave_type(code="ANNUAL", **over):
    fields = dict(name="استحقاقی", is_hourly=True)
    fields.update(over)
    lt, _ = LeaveType.objects.get_or_create(code=code, defaults=fields)
    return lt


def xlsx_bytes(rows: list[list], headers=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers or BALANCE_HEADER_ALIASES))
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TemplateTests(TestCase):
    def test_template_lists_live_leave_types(self):
        leave_type("ANNUAL")
        leave_type("SICK", name="استعلاجی", is_hourly=False)
        payload = build_balance_template()
        wb = load_workbook(io.BytesIO(payload), read_only=True)
        self.assertEqual(wb.sheetnames, ["مانده‌ها", "انواع مرخصی", "راهنما"])
        types_sheet = list(wb["انواع مرخصی"].iter_rows(values_only=True))
        codes = {row[0] for row in types_sheet[1:]}
        self.assertIn("ANNUAL", codes)
        self.assertIn("SICK", codes)
        main_header = list(wb["مانده‌ها"].iter_rows(values_only=True))[0]
        self.assertIn("personnel_code", main_header)
        wb.close()


class ImportBalanceTests(TestCase):
    def setUp(self):
        self.branch = make_branch()
        self.emp1, _ = make_employee(self.branch, code="12345")
        self.emp2, _ = make_employee(self.branch, code="12346")
        self.hr_emp, self.hr_user = make_employee(self.branch, code="HR-9")
        grant(self.hr_user, "HR")
        leave_type("ANNUAL")
        leave_type("SICK", name="استعلاجی")
        self.client = APIClient()

    def test_import_creates_entries_with_running_balance(self):
        data = xlsx_bytes([
            ["12345", "ANNUAL", "OPENING", 12480, "1404/01/01"],
            ["12346", "ANNUAL", "OPENING", 9480, "1404/01/01"],
            ["12345", "ANNUAL", "ADJUST", -480, "1404/05/01"],
        ])
        report = import_balances(data)
        self.assertEqual(report["total"], 3)
        self.assertEqual(report["created"], 3)
        self.assertEqual(report["skipped"], 0)
        self.assertFalse(report["dry_run"])
        annual = LeaveType.objects.get(code="ANNUAL")
        self.assertEqual(current_balance_minutes(self.emp1, annual), 12000)
        self.assertEqual(current_balance_minutes(self.emp2, annual), 9480)
        rows = LeaveLedger.objects.filter(employee=self.emp1).order_by("id")
        self.assertEqual([r.entry_kind for r in rows], ["OPENING", "ADJUST"])
        self.assertEqual([r.balance_after_minutes for r in rows], [12480, 12000])
        self.assertTrue(all(r.source_type == "HR_SEED_XLSX" for r in rows))

    def test_dry_run_touches_nothing(self):
        data = xlsx_bytes([["12345", "ANNUAL", "OPENING", 12480, "1404/01/01"]])
        report = import_balances(data, dry_run=True)
        self.assertEqual(report["created"], 1)
        self.assertTrue(report["dry_run"])
        self.assertEqual(LeaveLedger.objects.count(), 0)

    def test_bad_rows_reported_good_rows_land(self):
        data = xlsx_bytes([
            ["99999", "ANNUAL", "OPENING", 100, "1404/01/01"],   # پرسنلی ناشناس
            ["12345", "MYSTERY", "OPENING", 100, "1404/01/01"],  # نوع نامعتبر
            ["12345", "ANNUAL", "DEDUCTION", 100, "1404/01/01"],  # گردش غیرمجاز
            ["12345", "ANNUAL", "OPENING", 0, "1404/01/01"],      # صفر
            ["12345", "ANNUAL", "OPENING", 480, "1404/13/40"],    # تاریخ خراب
            ["12345", "ANNUAL", "OPENING", 480, "1404/01/01"],    # سالم
        ])
        report = import_balances(data)
        self.assertEqual(report["total"], 6)
        self.assertEqual(report["created"], 1)
        self.assertEqual(report["skipped"], 5)
        self.assertEqual(len(report["errors"]), 5)
        self.assertEqual(LeaveLedger.objects.count(), 1)

    def test_negative_minutes_only_for_adjust(self):
        data = xlsx_bytes([
            ["12345", "ANNUAL", "OPENING", -480, "1404/01/01"],   # منفی غیر ADJUST
            ["12345", "ANNUAL", "ADJUST", -480, "1404/01/01"],    # منفی ADJUST سالم
        ])
        report = import_balances(data)
        self.assertEqual(report["created"], 1)
        self.assertEqual(report["skipped"], 1)
        annual = LeaveType.objects.get(code="ANNUAL")
        self.assertEqual(current_balance_minutes(self.emp1, annual), -480)

    def test_persian_headers_accepted(self):
        data = xlsx_bytes(
            [["12345", "ANNUAL", 960, "1404/02/01"]],
            headers=["کد پرسنلی", "نوع مرخصی", "دقیقه", "تاریخ مبنا"],
        )
        report = import_balances(data)
        self.assertEqual(report["created"], 1)

    def test_missing_required_headers_raise(self):
        data = xlsx_bytes([["12345", "ANNUAL"]], headers=["کد پرسنلی", "نوع مرخصی"])
        with self.assertRaises(WorkbookError):
            import_balances(data)

    def test_garbage_bytes_raise_workbook_error(self):
        with self.assertRaises(WorkbookError):
            import_balances(b"not an xlsx at all")

    def test_empty_rows_skipped_silently(self):
        data = xlsx_bytes([
            [None, None, None, None, None],
            ["12345", "ANNUAL", "OPENING", 480, "1404/01/01"],
        ])
        report = import_balances(data)
        self.assertEqual(report["total"], 1)


class ImportApiTests(TestCase):
    def setUp(self):
        cache.clear()  # bucket throttle «import» per-IP است و بین تست‌ها می‌ماند
        self.branch = make_branch()
        self.emp, self.user = make_employee(self.branch, code="12345")
        self.hr_emp, self.hr_user = make_employee(self.branch, code="HR-9")
        grant(self.hr_user, "HR")
        leave_type("ANNUAL")
        self.client = APIClient()

    def _upload(self, payload: bytes, dry_run: bool = False):
        form = {
            "file": io.BytesIO(payload),
            "dry_run": "true" if dry_run else "false",
        }
        return self.client.post("/api/v1/leaves/ledger/import/", form, format="multipart")

    def test_hr_downloads_template(self):
        self.client.force_authenticate(self.hr_user)
        res = self.client.get("/api/v1/leaves/ledger/template/")
        self.assertEqual(res.status_code, 200)
        self.assertIn("spreadsheetml", res["Content-Type"])
        wb = load_workbook(io.BytesIO(res.content), read_only=True)
        self.assertIn("انواع مرخصی", wb.sheetnames)
        wb.close()

    def test_employee_cannot_touch_template_or_import(self):
        self.client.force_authenticate(self.user)
        res = self.client.get("/api/v1/leaves/ledger/template/")
        self.assertEqual(res.status_code, 403)
        res2 = self._upload(xlsx_bytes([["12345", "ANNUAL", "OPENING", 480, "1404/01/01"]]))
        self.assertEqual(res2.status_code, 403)

    def test_hr_import_round_trip(self):
        self.client.force_authenticate(self.hr_user)
        res = self._upload(
            xlsx_bytes([["12345", "ANNUAL", "OPENING", 12480, "1404/01/01"]]), dry_run=True
        )
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data["created"], 1)
        self.assertEqual(LeaveLedger.objects.count(), 0)

        res2 = self._upload(
            xlsx_bytes([["12345", "ANNUAL", "OPENING", 12480, "1404/01/01"]])
        )
        self.assertEqual(res2.status_code, 200)
        self.assertEqual(res2.data["created"], 1)
        self.assertEqual(LeaveLedger.objects.count(), 1)

    def test_invalid_file_is_400(self):
        self.client.force_authenticate(self.hr_user)
        res = self._upload(b"garbage")
        self.assertEqual(res.status_code, 400)

    def test_missing_file_is_400(self):
        self.client.force_authenticate(self.hr_user)
        res = self.client.post("/api/v1/leaves/ledger/import/", {"dry_run": "true"})
        self.assertEqual(res.status_code, 400)
