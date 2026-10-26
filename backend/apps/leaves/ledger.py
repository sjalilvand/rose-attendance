"""LeaveLedger activation — the append-only balance stream goes live (§11).

Until this module the ledger table existed but stayed silent (see the old
services.py note: opening balances awaited HR input). It is now wired
end-to-end:

- HR seeds OPENING / ACCRUAL / ADJUST entries — bulk JSON with dry-run
  (leaves/api.LeaveLedgerViewSet.seed); the stream stays append-only, so
  corrections are *new* entries (ADJUST), never edits or deletes
- approving a LEAVE segment appends a DEDUCTION (−computed_minutes) atomically
  with the decision (hooked in services.decide_segment)
- a leave type becomes "tracked" for an employee once its first entry exists;
  untracked (employee, type) pairs never block submission — sick leave etc.
  can stay unmanaged
- the authoritative balance is SUM(minutes) of the stream (race-safe under
  concurrent decisions); the stored balance_after_minutes is a running
  display convenience, recomputed from the current sum at each append

Restores (REJECT after approve) don't exist yet: segments are decided once
and cancellation only works on PENDING, so there is no approval-reversal
path. If HR voids a wrong approval, a positive ADJUST entry is the answer.
"""
from __future__ import annotations

import datetime as dt
import io

import jdatetime
from django.db import transaction
from django.db.models import Max, Sum
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.api_utils import hm_minutes
from apps.core.enums import ApprovalAction, LeaveEntryKind
from apps.people.models import Employee

from .models import LeaveLedger, LeaveType

SEED_KINDS = (LeaveEntryKind.OPENING, LeaveEntryKind.ACCRUAL, LeaveEntryKind.ADJUST)
MAX_SEED_ROWS = 10000


class LedgerError(Exception):
    """Domain error surfaced as a coded API response (mirrors RequestError)."""

    def __init__(self, message: str, code: str = "INVALID"):
        super().__init__(message)
        self.message = message
        self.code = code


def parse_basis_date(value) -> dt.date | None:
    """Jalali `1404-01-01` / `1404/01/01` first, then ISO — HR speaks Jalali."""
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value).strip()
    normalized = text.replace("/", "-")
    parts = normalized.split("-")
    if len(parts) == 3:
        try:
            jy, jm, jd = (int(p) for p in parts)
            if 1300 <= jy <= 1500:  # شمسی
                return jdatetime.date(jy, jm, jd).togregorian()
        except ValueError:
            pass
    try:
        return dt.date.fromisoformat(normalized)
    except ValueError as exc:
        raise LedgerError(f"تاریخ نامعتبر است: {text}") from exc


def current_balance_minutes(employee, leave_type) -> int:
    row = LeaveLedger.objects.filter(employee=employee, leave_type=leave_type).aggregate(
        total=Sum("minutes")
    )
    return int(row["total"] or 0)


def is_tracked(employee, leave_type) -> bool:
    return LeaveLedger.objects.filter(employee=employee, leave_type=leave_type).exists()


def append_entry(
    *,
    employee,
    leave_type,
    entry_kind: str,
    minutes: int,
    basis_date: dt.date,
    source_type: str = "",
    source_id: int | None = None,
) -> LeaveLedger:
    """Single append with running balance — used by decisions and adjust."""
    with transaction.atomic():
        base = current_balance_minutes(employee, leave_type)
        return LeaveLedger.objects.create(
            employee=employee,
            leave_type=leave_type,
            entry_kind=entry_kind,
            minutes=minutes,
            balance_after_minutes=base + minutes,
            basis_date=basis_date,
            source_type=source_type or "",
            source_id=source_id,
        )


def enforce_sufficient_balance(employee, leave_type, minutes: int) -> None:
    """Submission-time guard — raises LedgerError(INSUFFICIENT_BALANCE).

    Silent pass for untracked (employee, type) pairs: seeding is opt-in per
    type, so HR chooses which types carry numeric balances (ANNUAL typically)
    and which don't (SICK, UNPAID…).
    """
    if not is_tracked(employee, leave_type):
        return
    balance = current_balance_minutes(employee, leave_type)
    if balance < minutes:
        raise LedgerError(
            f"ماندهٔ {leave_type.name} کافی نیست (ماندهٔ فعلی: {hm_minutes(balance)}).",
            "INSUFFICIENT_BALANCE",
        )


def apply_segment_decision(segment, action: str) -> None:
    """Ledger side-effect of a segment decision — call inside the same tx.

    Only LEAVE × APPROVE touches the stream: −computed_minutes, anchored to
    the segment's start date and traced back to RequestSegment.<id>.
    """
    if action != ApprovalAction.APPROVE or segment.request_kind != "LEAVE":
        return
    leave_request = segment.leave_request
    append_entry(
        employee=leave_request.employee,
        leave_type=leave_request.leave_type,
        entry_kind=LeaveEntryKind.DEDUCTION,
        minutes=-int(segment.computed_minutes),
        basis_date=timezone.localtime(segment.starts_at).date(),
        source_type="RequestSegment",
        source_id=segment.id,
    )


def balances_for(employee) -> list[dict]:
    """Per-active-type balances for one employee (tracked ones carry sums)."""
    from django.db.models import Count

    rows = (
        LeaveLedger.objects.filter(employee=employee)
        .values("leave_type_id")
        .annotate(total=Sum("minutes"), entries=Count("id"), last_basis=Max("basis_date"))
    )
    by_type = {row["leave_type_id"]: row for row in rows}

    out: list[dict] = []
    for leave_type in LeaveType.objects.filter(is_active=True).order_by("id"):
        row = by_type.get(leave_type.id)
        out.append(
            {
                "leave_type": leave_type.code,
                "leave_type_name": leave_type.name,
                "is_hourly": leave_type.is_hourly,
                "tracked": row is not None,
                "balance_minutes": int(row["total"] or 0) if row else None,
                "entries": row["entries"] if row else 0,
                "last_basis_date": row["last_basis"] if row else None,
            }
        )
    return out


# ---------------------------------------------------------------------------
# Excel workbook — dedicated balance template + importer (فاز ۲، گام ۶)
# ---------------------------------------------------------------------------
BALANCE_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "personnel_code": ("personnel_code", "کد پرسنلی"),
    "leave_type": ("leave_type", "کد نوع مرخصی", "نوع مرخصی"),
    "entry_kind": ("entry_kind", "نوع گردش"),
    "minutes": ("minutes", "دقیقه", "مانده (دقیقه)"),
    "basis_date": ("basis_date", "تاریخ مبنا", "تاریخ مبنای شمسی"),
}
BALANCE_REQUIRED = ("personnel_code", "leave_type", "minutes")
BALANCE_SAMPLE_ROW = ["12345", "ANNUAL", "OPENING", 12480, "1404/01/01"]
MAX_REPORT_ERRORS = 50


class WorkbookError(Exception):
    """Workbook-level problem (unreadable file, missing required headers)."""


def build_balance_template() -> bytes:
    """Xlsx template — «مانده‌ها» + dynamic «انواع مرخصی» reference + «راهنما».

    The leave-type sheet is generated from live active types so HR copies the
    codes instead of typing them; unknown codes are the #1 import error.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "مانده‌ها"
    ws.append(list(BALANCE_HEADER_ALIASES))
    ws.append(BALANCE_SAMPLE_ROW)
    for col, width in enumerate((14, 16, 12, 14, 16), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws2 = wb.create_sheet("انواع مرخصی")
    ws2.append(["کد", "نام", "ساعتی؟", "فعال"])
    for t in LeaveType.objects.filter(is_active=True).order_by("id"):
        ws2.append([t.code, t.name, "بله" if t.is_hourly else "خیر", "فعال"])
    for col, width in zip("ABCD", (16, 30, 10, 10)):
        ws2.column_dimensions[col].width = width
    ws2["A1"].font = Font(bold=True)

    ws3 = wb.create_sheet("راهنما")
    guide = [
        ("personnel_code", "الزامی", "کد پرسنلی واردشده در سامانه (ستون را متنی نگه دارید)"),
        ("leave_type", "الزامی", "کد نوع مرخصی فعال — فهرست در برگهٔ «انواع مرخصی»"),
        ("entry_kind", "اختیاری", "OPENING / ACCRUAL / ADJUST — پیش‌فرض OPENING؛ مقدار منفی فقط برای ADJUST"),
        ("minutes", "الزامی", "دقیقهٔ صحیح مخالف صفر — مثال: 12480 = ۲۶ روز ۸ ساعته؛ 960 = ۲ روز"),
        ("basis_date", "اختیاری", "تاریخ شمسی 1404-01-01 یا 1404/01/01 (یا میلادی ISO) — پیش‌فرض امروز"),
        ("", "", "دفتر مانده تغییرناپذیر است: هر ردیف یک گردش جدید ثبت می‌کند، ردیف تکراری = گردش دوباره"),
        ("", "", "قبل از ثبت نهایی، همان فایل را با «بررسی بدون ثبت» اعتبارسنجی کنید"),
    ]
    ws3.append(["ستون", "الزام", "توضیح"])
    for row in guide:
        ws3.append(list(row))
    for col, width in zip("ABC", (16, 10, 84)):
        ws3.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def import_balances(data: bytes, *, dry_run: bool = False) -> dict:
    """Append ledger entries from an xlsx workbook; row-level error report.

    Mirrors the JSON seed (LeaveLedgerViewSet.seed) semantics exactly:
    append-only — every valid row becomes a new OPENING/ACCRUAL/ADJUST entry,
    running balance_after recomputed via one grouped query. Good rows land,
    bad rows are reported (cap MAX_REPORT_ERRORS = 50) and skipped.
    """
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a zoo of exception types
        raise WorkbookError("فایل قابل خواندن نیست؛ قالب xlsx الزامی است.") from exc

    errors: list[dict] = []
    total_rows = 0
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        try:
            header = list(next(rows))
        except StopIteration as exc:
            raise WorkbookError("کاربرگ خالی است.") from exc

        canonical: dict[str, int] = {}
        for idx, raw in enumerate(header):
            text = _cell_text(raw)
            if text is None:
                continue
            key = text.lower()
            for canon, aliases in BALANCE_HEADER_ALIASES.items():
                if key in aliases and canon not in canonical:
                    canonical[canon] = idx
        missing = [name for name in BALANCE_REQUIRED if name not in canonical]
        if missing:
            raise WorkbookError("ستون‌های الزامی یافت نشد: " + "، ".join(missing))

        wanted_codes = set()
        probe = list(rows)
        for raw_row in probe:
            row = {
                name: raw_row[idx] if idx < len(raw_row) else None
                for name, idx in canonical.items()
            }
            code = _cell_text(row.get("personnel_code"))
            if code:
                wanted_codes.add(code)
        employee_map = {e.personnel_code: e for e in Employee.objects.filter(personnel_code__in=wanted_codes)}
        type_map = {t.code: t for t in LeaveType.objects.filter(is_active=True)}
        kind_values = tuple(kind for kind, _label in LeaveEntryKind.choices)

        parsed: list[dict] = []
        for n, raw_row in enumerate(probe, start=2):  # header is row 1
            row = {
                name: raw_row[idx] if idx < len(raw_row) else None
                for name, idx in canonical.items()
            }
            if all(value in (None, "") for value in row.values()):
                continue  # fully empty line
            total_rows += 1

            code = _cell_text(row.get("personnel_code"))
            employee = employee_map.get(code or "")
            if employee is None:
                errors.append({"row": n, "personnel_code": code or "", "message": "کد پرسنلی یافت نشد."})
                continue

            type_code = _cell_text(row.get("leave_type")) or ""
            leave_type = type_map.get(type_code)
            if leave_type is None:
                errors.append(
                    {"row": n, "personnel_code": code,
                     "message": f"نوع مرخصی فعال «{type_code}» یافت نشد (برگهٔ «انواع مرخصی»)."}
                )
                continue

            entry_kind = (_cell_text(row.get("entry_kind")) or LeaveEntryKind.OPENING).upper()
            if entry_kind not in kind_values:
                errors.append(
                    {"row": n, "personnel_code": code,
                     "message": f"نوع گردش نامعتبر: {entry_kind} — یکی از {', '.join(kind_values)}."}
                )
                continue
            if entry_kind not in SEED_KINDS:
                errors.append(
                    {"row": n, "personnel_code": code,
                     "message": f"نوع گردش «{entry_kind}» از طریق ورود فایل مجاز نیست (فقط {', '.join(SEED_KINDS)})."}
                )
                continue

            minutes_raw = row.get("minutes")
            try:
                minutes = int(minutes_raw)
            except (TypeError, ValueError):
                errors.append({"row": n, "personnel_code": code, "message": "دقیقه باید عدد صحیح باشد."})
                continue
            if minutes == 0 or (minutes < 0 and entry_kind != LeaveEntryKind.ADJUST):
                errors.append(
                    {"row": n, "personnel_code": code,
                     "message": "دقیقه باید مخالف صفر باشد؛ مقدار منفی فقط برای ADJUST."}
                )
                continue

            try:
                basis_date = parse_basis_date(row.get("basis_date")) or timezone.localdate()
            except LedgerError as exc:
                errors.append({"row": n, "personnel_code": code, "message": exc.message})
                continue

            parsed.append(
                {"row": n, "employee": employee, "leave_type": leave_type,
                 "entry_kind": entry_kind, "minutes": minutes, "basis_date": basis_date}
            )

        # Running balances per (employee, type) — one grouped query, then in-memory.
        running: dict[tuple[int, int], int] = {}
        if parsed:
            emp_ids = {r["employee"].id for r in parsed}
            type_ids = {r["leave_type"].id for r in parsed}
            sums = (
                LeaveLedger.objects.filter(employee_id__in=emp_ids, leave_type_id__in=type_ids)
                .values("employee_id", "leave_type_id")
                .annotate(total=Sum("minutes"))
            )
            running = {(s["employee_id"], s["leave_type_id"]): int(s["total"] or 0) for s in sums}

        to_create: list[LeaveLedger] = []
        for r in parsed:
            key = (r["employee"].id, r["leave_type"].id)
            running[key] = running.get(key, 0) + r["minutes"]
            if not dry_run:
                to_create.append(
                    LeaveLedger(
                        employee=r["employee"],
                        leave_type=r["leave_type"],
                        entry_kind=r["entry_kind"],
                        minutes=r["minutes"],
                        balance_after_minutes=running[key],
                        basis_date=r["basis_date"],
                        source_type="HR_SEED_XLSX",
                    )
                )
        if not dry_run and to_create:
            with transaction.atomic():
                LeaveLedger.objects.bulk_create(to_create, batch_size=500)
    finally:
        wb.close()

    return {
        "total": total_rows,
        "created": len(parsed) if dry_run else len(to_create),
        "updated": 0,
        "skipped": len(errors),
        "errors": errors[:MAX_REPORT_ERRORS],
        "dry_run": dry_run,
    }
